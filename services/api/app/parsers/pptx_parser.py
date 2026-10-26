# ================================================================
# services/api/app/parsers/pptx_parser.py
# ================================================================
# Extracts text from PowerPoint (.pptx) files, one slide per page.
#
# WHY python-pptx:
# Reads the slide XML directly, so it needs no PowerPoint or
# LibreOffice installation. Works the same on Windows and Linux (CI).
# ================================================================

from io import BytesIO
from pptx import Presentation

from services.api.app.parsers.pdf_parser import ParsedDocument


def parse_pptx(file_bytes: bytes) -> ParsedDocument:
    """
    Extracts text from every slide. Each slide becomes one page,
    so retrieval results can point to a slide number.
    Slide text comes from shapes with text frames and from tables.
    """
    try:
        presentation = Presentation(BytesIO(file_bytes))
    except Exception as e:
        raise ValueError(f"Failed to open PowerPoint file: {e}")

    pages = []
    for slide in presentation.slides:
        parts = []
        for shape in slide.shapes:
            # Plain text boxes and placeholders
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
            # Tables inside the slide, flattened row by row
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        pages.append("\n".join(parts))

    # Keep only non-empty slides for the full text, but keep the
    # slide list complete so page numbers still match slide numbers
    full_text = "\n\n".join(p for p in pages if p.strip())
    if not full_text.strip():
        raise ValueError("No extractable text found in PowerPoint file")

    return ParsedDocument(
        full_text=full_text,
        pages=pages,
        page_count=len(pages),
        metadata={"slide_count": len(pages)},
        parser_used="python-pptx",
    )