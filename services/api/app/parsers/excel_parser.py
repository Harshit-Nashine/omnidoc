# ================================================================
# services/api/app/parsers/excel_parser.py
# ================================================================
# Extracts text from Excel (.xlsx) workbooks, one sheet per page.
#
# WHY openpyxl:
# Reads .xlsx directly with no Excel installation needed. Using
# read_only=True streams the sheet instead of loading it all into
# memory, which matters for large workbooks.
# ================================================================

from io import BytesIO
from openpyxl import load_workbook

from services.api.app.parsers.pdf_parser import ParsedDocument


def parse_excel(file_bytes: bytes) -> ParsedDocument:
    """
    Extracts text from every sheet. Each row becomes one line with
    cells separated by " | ", so the row context survives chunking.
    Formula cells are read as their cached values (data_only=True).
    """
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"Failed to open Excel file: {e}")

    pages = []
    for sheet in workbook.worksheets:
        lines = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if cells:
                lines.append(" | ".join(cells))
        pages.append(f"Sheet: {sheet.title}\n" + "\n".join(lines))
    workbook.close()

    full_text = "\n\n".join(pages)
    if not any(line.strip() for p in pages for line in p.splitlines()[1:]):
        raise ValueError("No extractable data found in Excel file")

    return ParsedDocument(
        full_text=full_text,
        pages=pages,
        page_count=len(pages),
        metadata={"sheet_count": len(pages)},
        parser_used="openpyxl",
    )