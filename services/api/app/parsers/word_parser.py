# ================================================================
# services/api/app/parsers/word_parser.py
# ================================================================

from docx import Document
from io import BytesIO


def parse_word_document(file_content: bytes):
    """
    Extracts text from a Word document (.docx).
    Returns a ParsedDocument matching the same structure
    as pdf_parser and image_parser.
    """
    # Import here to use the same dataclass as other parsers
    from services.api.app.parsers.pdf_parser import ParsedDocument

    try:
        doc = Document(BytesIO(file_content))
    except Exception as e:
        raise ValueError(f"Failed to parse Word document: {str(e)}")

    # Extract paragraphs
    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)

    # Extract tables
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells]
            row_text = " | ".join(cell for cell in row_cells if cell)
            if row_text.strip():
                paragraphs.append(row_text)

    full_text = "\n\n".join(paragraphs)

    if not full_text.strip():
        raise ValueError("No extractable text found in Word document")

    # Word has no page concept — treat as one page
    pages = [full_text]

    return ParsedDocument(
        full_text=full_text,
        pages=pages,
        page_count=1,
        metadata={"source": "word", "paragraph_count": len(paragraphs)},
        parser_used="python-docx",
    )